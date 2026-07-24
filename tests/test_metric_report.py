import json
import math
import tempfile
import unittest
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from agents.pid_qlearning_agent import PIDQLearningAgent
from controllers.controller_manager import ControllerManager
from dynamic_system.cart_pole_dynamic_system import CartPoleDynamicSystem
from main import build_capacities, load_config
from metrics.metric_collector import MetricCollector
from metrics.metric_processing import MetricProcessor
from metrics.metric_report import MetricReport
from result_handler import ResultHandler
from rewards.reward_calculator import RewardCalculator
from runtime_arrays import build_runtime_arrays
from simulation_manager import SimulationManager


class MetricReportScientificTests(unittest.TestCase):
    def build_components(self, config, output_dir, run_id="test_run"):
        step_capacity, interval_capacity = build_capacities(config)
        runtime = build_runtime_arrays(step_capacity, interval_capacity, config)
        dynamic_system = CartPoleDynamicSystem(runtime, config["dynamic_system"])
        controller_manager = ControllerManager(
            runtime,
            config["control"],
            config["agent"],
            config["simulation"]["dt_sec"],
        )
        agent = PIDQLearningAgent(runtime, config["agent"], config["control"])
        metric_report = MetricReport(runtime, config, agent, run_id)
        result_handler = ResultHandler(
            output_dir,
            config["data"],
            config["termination"],
            metric_report,
        )
        metric_processor = MetricProcessor(runtime, config["metrics"], config["simulation"]["dt_sec"])
        reward_calculator = RewardCalculator(runtime, config["reward"], config["agent"], config["termination"])
        metric_collector = MetricCollector(runtime, config["data"], result_handler)
        simulation_manager = SimulationManager(
            runtime=runtime,
            config=config,
            dynamic_system=dynamic_system,
            controller_manager=controller_manager,
            agent=agent,
            metric_processor=metric_processor,
            metric_report=metric_report,
            reward_calculator=reward_calculator,
            metric_collector=metric_collector,
            result_handler=result_handler,
        )
        return runtime, agent, metric_report, result_handler, simulation_manager

    def test_scientific_reductions_and_denominators(self):
        config = load_config("config/config.yaml")
        config["simulation"]["dt_sec"] = 0.1
        config["simulation"]["decision_interval_sec"] = 0.2
        config["simulation"]["episode_duration_sec"] = 0.4
        config["termination"]["stabilization"]["persistence_sec"] = 0.2
        with tempfile.TemporaryDirectory() as output_dir:
            runtime, agent, metric_report, _, _ = self.build_components(config, output_dir)
            agent.reset_episode()
            n_steps = 4
            n_intervals = 3

            runtime.step.t_sec[:5] = [0.0, 0.1, 0.2, 0.3, 0.4]
            runtime.step.pendulum_angle_raw[:5] = [0.1, 0.03, 0.0, 0.0, 0.03]
            runtime.step.pendulum_velocity_raw[:5] = 0.0
            runtime.step.cart_position_raw[:5] = 0.0
            runtime.step.cart_velocity_raw[:5] = 0.0
            runtime.step.error_pendulum_angle_raw[:4] = [-1.0, 0.0, 1.0, 2.0]
            runtime.step.error_cart_position_raw[:4] = [0.0, 1.0, 0.0, -1.0]
            runtime.step.u_demand_pendulum_angle_norm[:4] = [1.0, -1.0, 1.0, -1.0]
            runtime.step.u_demand_cart_position_norm[:4] = [-1.0, 1.0, -1.0, 1.0]
            runtime.step.u_total_raw[:4] = 0.0
            runtime.step.u_total[:4] = [0.0, 0.5, -0.5, 0.0]
            runtime.step.is_saturated_global[:4] = 0.0
            runtime.step.cart_force[1:5] = 1.0

            runtime.interval.local_health_01_pendulum_angle[:3] = [0.5, 0.8, 0.9]
            runtime.interval.local_health_01_cart_position[:3] = [0.5, 0.4, 0.5]
            runtime.interval.cooperative_weight_effective_pendulum_angle[:3] = 0.5
            runtime.interval.cooperative_weight_effective_cart_position[:3] = 0.5
            runtime.interval.cooperative_reward_signed[:3] = [-1.0, 0.0, 1.0]
            runtime.interval.local_reward_signed_pendulum_angle[:3] = [-0.5, 0.0, 0.5]
            runtime.interval.local_reward_signed_cart_position[:3] = [0.5, 0.0, -0.5]
            for gain_agent_config in config["agent"]["gain_agents"]:
                agent_id = gain_agent_config["id"]
                getattr(runtime.interval, f"gain_value_{agent_id}")[:3] = 1.0

            success_code = config["termination"]["reason_codes"]["stabilization_success"]
            metric_report.compute_episode(
                episode_idx=0,
                n_steps=n_steps,
                n_intervals=n_intervals,
                termination_reason_code=success_code,
                training_enabled=False,
            )

            self.assertAlmostEqual(runtime.episode.iae_pendulum_angle, 0.4)
            self.assertAlmostEqual(runtime.episode.ise_pendulum_angle, 0.6)
            self.assertAlmostEqual(runtime.episode.rms_error_pendulum_angle, math.sqrt(1.5))
            self.assertAlmostEqual(runtime.episode.stabilization_start_sec, 0.2)
            self.assertAlmostEqual(runtime.episode.stable_block_max_sec, 0.2)
            self.assertAlmostEqual(runtime.episode.stable_occupancy, 0.5)
            self.assertEqual(runtime.episode.joint_stability_loss_count, 1)
            self.assertAlmostEqual(runtime.episode.joint_post_entry_duration_sec, 0.2)
            self.assertAlmostEqual(runtime.episode.joint_stability_loss_rate_hz, 5.0)
            self.assertAlmostEqual(runtime.episode.cancellation_denominator, 0.8)
            self.assertAlmostEqual(runtime.episode.cancellation_index, 1.0)
            self.assertEqual(runtime.episode.health_transition_count, 2)
            self.assertAlmostEqual(runtime.episode.concurrent_progress_frequency, 0.5)
            self.assertAlmostEqual(runtime.episode.deterioration_frequency, 0.5)
            self.assertAlmostEqual(runtime.episode.compensatory_transition_frequency, 0.5)
            self.assertAlmostEqual(runtime.episode.cooperative_reward_positive_proportion, 1.0 / 3.0)
            self.assertAlmostEqual(runtime.episode.cooperative_reward_zero_proportion, 1.0 / 3.0)
            self.assertAlmostEqual(runtime.episode.cooperative_reward_negative_proportion, 1.0 / 3.0)
            self.assertEqual(runtime.episode.steady_state_window_applicable_pendulum_angle, 0)
            self.assertIsNone(runtime.episode.steady_state_error_signed_pendulum_angle)
            self.assertEqual(runtime.episode.learning_metrics_applicable_kp_pendulum_angle, 0)
            self.assertIsNone(runtime.episode.q_delta_rms_kp_pendulum_angle)

    def test_agent_exposes_q_and_policy_reductions(self):
        config = load_config("config/config.yaml")
        step_capacity, interval_capacity = build_capacities(config)
        runtime = build_runtime_arrays(step_capacity, interval_capacity, config)
        agent = PIDQLearningAgent(runtime, config["agent"], config["control"])
        agent.reset_episode()
        track = agent.tracks[0]
        state_idx = track.current_state_idx
        track.q_table[state_idx, 0] = 1.0
        metrics = {item["agent_id"]: item for item in agent.episode_learning_metrics()}
        first_agent = metrics[track.agent_id]
        self.assertAlmostEqual(first_agent["q_delta_rms"], 1.0 / math.sqrt(track.q_table.size))
        self.assertEqual(first_agent["visited_state_count"], 1)
        self.assertEqual(first_agent["policy_stability"], 0.0)

    def test_short_run_persists_canonical_workbook(self):
        config = load_config("config/config.yaml")
        config["simulation"]["n_episodes"] = 2
        config["simulation"]["episode_duration_sec"] = 0.04
        config["simulation"]["dt_sec"] = 0.01
        config["simulation"]["decision_interval_sec"] = 0.02
        config["data"]["episodes_per_dataset_chunk"] = 2
        config["data"]["save_agent_state"] = False
        config["visualization"]["enabled"] = False

        with tempfile.TemporaryDirectory() as output_dir:
            _, _, metric_report, result_handler, simulation_manager = self.build_components(
                config,
                output_dir,
                run_id="short_run",
            )
            result_handler.save_metadata({"run_id": "short_run", "config": config})
            simulation_manager.run()

            workbook_path = Path(output_dir) / "metric_report.xlsx"
            workbook = load_workbook(workbook_path, data_only=False, read_only=False)
            self.assertEqual(workbook.sheetnames, ["Episode_Metrics", "Metric_Catalog", "Run_Config"])
            self.assertEqual(workbook["Episode_Metrics"].max_row, 3)
            self.assertEqual(workbook["Episode_Metrics"].max_column, len(metric_report.episode_columns))
            self.assertEqual(workbook["Metric_Catalog"].max_row, len(metric_report.catalog_rows) + 1)
            self.assertEqual(workbook["Run_Config"].max_row, 2)
            self.assertEqual(workbook["Run_Config"]["A2"].value, "short_run")
            self.assertTrue((Path(output_dir) / "summary.xlsx").exists())

            chunk_path = Path(output_dir) / "dataset_chunks" / "dataset_chunk_000000.json"
            chunk = json.loads(chunk_path.read_text(encoding="utf-8"))
            self.assertEqual(chunk["n_episodes"], 2)
            end_data = chunk["episodes"][0]["end_episode_data"]
            self.assertEqual(end_data["failure"], 1)
            self.assertIn("q_delta_rms_kp_pendulum_angle", end_data)
            self.assertEqual(set(metric_report.metric_signal_names) - set(end_data), set())


if __name__ == "__main__":
    unittest.main()
